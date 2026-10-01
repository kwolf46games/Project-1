"""Question/response banks: one Markdown file per bank in banks/.

File format (hand-editable):

    # Bank title                                  (optional)

    ## How would you describe your leadership style?
    also: how do you lead a team | what kind of leader are you
    follows: Tell me about yourself               (optional, see below)
    tags: leadership, behavioral                  (optional)
    skeleton: one-line outline shown above the answer   (optional)

    The response goes here, in plain text.
    **Bold** marks a line to land, {{double braces}} mark a placeholder to fix.

Metadata lines (also / follows / tags / skeleton) sit directly under the "## question"
heading; everything after them is the response body.

Follow-up questions
    "follows:" says this question is a follow-up to another one in the same (or another
    active) bank, by that question's exact text; separate several with "|".  After the
    parent's answer is shown, the overlay lists its follow-ups and listens for them with
    extra confidence.  Tag a question "followup" (no "follows:" needed) for generic
    follow-ups such as "Can you give me an example?" that can come after any answer.
"""
from __future__ import annotations

import csv
import json
import re
from dataclasses import dataclass, field
from pathlib import Path

ROOT = Path(__file__).resolve().parent
BANKS_DIR = ROOT / "banks"
CONFIG_PATH = ROOT / "config.json"

DEFAULT_CONFIG = {
    "active_banks": [],
    "whisper_model": "base.en",
    "match_threshold": 0.78,
    "silence_seconds": 0.7,
    "font_size": 13,
    "opacity": 0.93,
    "whisper_prompt": "",
    "output_device": "",
    "followups": True,               # recognise follow-up questions using the answer just shown
    "followup_window_seconds": 120,  # how long an answer stays "the topic" for follow-ups
    "early_silence_seconds": 0.35,   # start transcribing after this much quiet (0 = wait the full silence_seconds)
    "reconnect_idle_seconds": 45,    # re-open the audio device after this long with no sound (0 = off)
    "generate": True,                # draft a spoken answer for each question (needs the anthropic package + a key)
    "generate_model": "claude-opus-5-5",
    "generate_effort": "low",        # low = fastest; medium/high think longer
    "generate_words": 130,           # about how long a drafted answer should be
    "generate_matches": 3,           # how many of your prepared answers are sent as the source material
    "generate_max_tokens": 4000,
}

META_KEYS = ("also", "follows", "tags", "skeleton")
FOLLOWUP_TAGS = {"followup", "followups", "follow-up", "follow-ups", "follow up", "follow ups"}


@dataclass
class Entry:
    question: str
    response: str = ""
    also: list[str] = field(default_factory=list)
    tags: list[str] = field(default_factory=list)
    skeleton: str = ""
    bank: str = ""
    follows: list[str] = field(default_factory=list)   # question text(s) this is a follow-up to

    def phrasings(self) -> list[str]:
        return [self.question, *self.also]

    @property
    def is_generic_followup(self) -> bool:
        """Tagged "followup": may be asked after any answer ("Can you give me an example?")."""
        return any(t.strip().lower() in FOLLOWUP_TAGS for t in self.tags)

    def to_markdown(self) -> str:
        lines = [f"## {self.question}"]
        if self.also:
            lines.append("also: " + " | ".join(self.also))
        if self.follows:
            lines.append("follows: " + " | ".join(self.follows))
        if self.tags:
            lines.append("tags: " + ", ".join(self.tags))
        if self.skeleton:
            lines.append("skeleton: " + self.skeleton)
        lines += ["", self.response.strip(), ""]
        return "\n".join(lines)


@dataclass
class Bank:
    name: str
    title: str = ""
    entries: list[Entry] = field(default_factory=list)

    @property
    def path(self) -> Path:
        return BANKS_DIR / f"{self.name}.md"

    def save(self) -> None:
        BANKS_DIR.mkdir(exist_ok=True)
        head = f"# {self.title}\n\n" if self.title else ""
        _write_atomic(self.path, head + "\n".join(e.to_markdown() for e in self.entries))

    def find(self, query: str) -> Entry:
        """Look up an entry by 1-based index or case-insensitive question substring."""
        if query.isdigit():
            i = int(query) - 1
            if 0 <= i < len(self.entries):
                return self.entries[i]
            raise KeyError(f"No question #{query} in '{self.name}' (it has {len(self.entries)}).")
        hits = [e for e in self.entries if query.lower() in e.question.lower()]
        if len(hits) == 1:
            return hits[0]
        if not hits:
            raise KeyError(f"No question matching '{query}' in '{self.name}'.")
        raise KeyError(f"'{query}' matches {len(hits)} questions; be more specific or use the number.")

    def rename_question(self, old: str, new: str) -> None:
        """Keep "follows:" links working when a question's wording changes."""
        for e in self.entries:
            e.follows = [new if f.strip().lower() == old.strip().lower() else f for f in e.follows]


def _write_atomic(path: Path, text: str) -> None:
    """Write via a temp file + rename so readers never see a half-written file."""
    tmp = path.with_name(path.name + ".tmp")
    tmp.write_text(text, encoding="utf-8")
    tmp.replace(path)


# ---------- parsing ----------

def parse_markdown(text: str, name: str) -> Bank:
    bank = Bank(name=name)
    chunks = re.split(r"^## +", text, flags=re.M)
    title_m = re.search(r"^# +(.+)$", chunks[0], flags=re.M)
    if title_m:
        bank.title = title_m.group(1).strip()
    for chunk in chunks[1:]:
        lines = chunk.splitlines()
        entry = Entry(question=lines[0].strip(), bank=name)
        i = 1
        while i < len(lines):
            m = re.match(r"^(also|follows|tags|skeleton):\s*(.*)$", lines[i], flags=re.I)
            if not m:
                break
            key, val = m.group(1).lower(), m.group(2).strip()
            if key == "also":
                entry.also = [p.strip() for p in val.split("|") if p.strip()]
            elif key == "follows":
                entry.follows = [p.strip() for p in val.split("|") if p.strip()]
            elif key == "tags":
                entry.tags = [t.strip() for t in val.split(",") if t.strip()]
            else:
                entry.skeleton = val
            i += 1
        entry.response = "\n".join(lines[i:]).strip()
        bank.entries.append(entry)
    return bank


IMPORT_TYPES = (".docx", ".xlsx", ".xlsm", ".csv", ".md", ".txt")

COLUMN_NAMES = {
    "question": ("question", "questions", "q", "prompt", "interview question"),
    "response": ("response", "responses", "answer", "answers", "a", "script", "my answer", "sample answer"),
    "also": ("also", "other phrasings", "phrasings", "alternates", "variations"),
    "follows": ("follows", "follow-up to", "follow up to", "followup to", "follows up", "parent", "parent question"),
    "tags": ("tags", "tag", "category", "type", "topic"),
    "skeleton": ("skeleton", "outline", "summary", "key points"),
}
QUESTION_WORDS = ("what", "why", "how", "when", "where", "which", "who", "tell me", "walk me", "describe",
                  "give me", "can you", "could you", "do you", "have you", "are you", "would you", "talk about")
NUMBER_PREFIX = re.compile(r"^\s*(?:(?:Q(?:uestion)?|[A-Z])\s*\d+[.):\-]?|\d+[.):\-])\s*", re.I)


def _looks_like_question(text: str) -> bool:
    t = _clean_question(text).lower().strip("“”\"' ")
    return t.endswith("?") or t.startswith(QUESTION_WORDS)


def _clean_question(text: str) -> str:
    return NUMBER_PREFIX.sub("", text.strip().strip("*")).strip()


def _rows_to_bank(rows: list[list[str]], name: str) -> Bank:
    """Spreadsheet rows -> bank. Finds columns by header name; with no recognizable header,
    uses the first column as the question and the second as the response."""
    rows = [[("" if c is None else str(c)).strip() for c in r] for r in rows if any(c for c in r if c)]
    if not rows:
        return Bank(name=name)
    header = [h.lower() for h in rows[0]]
    cols = {key: next((i for i, h in enumerate(header) if h in names), None) for key, names in COLUMN_NAMES.items()}
    if cols["question"] is None:
        cols = {"question": 0, "response": 1 if len(rows[0]) > 1 else None, "also": None, "follows": None,
                "tags": None, "skeleton": None}
    else:
        rows = rows[1:]
    if cols["response"] is None:
        cols["response"] = next((i for i in range(len(header)) if i not in cols.values()), None)

    def get(row, key):
        i = cols.get(key)
        return row[i] if i is not None and i < len(row) else ""

    bank = Bank(name=name)
    for row in rows:
        q = _clean_question(get(row, "question"))
        if q:
            bank.entries.append(Entry(
                question=q, response=get(row, "response"),
                also=[p.strip() for p in re.split(r"[|\n]", get(row, "also")) if p.strip()],
                follows=[p.strip() for p in re.split(r"[|\n]", get(row, "follows")) if p.strip()],
                tags=[t.strip() for t in get(row, "tags").split(",") if t.strip()],
                skeleton=get(row, "skeleton"), bank=name,
            ))
    return bank


def _blocks_to_bank(blocks: list[tuple[int | None, str]], name: str) -> Bank:
    """Generic document -> bank. `blocks` are (heading level or None, text) in reading order.

    Tries, in order: "Q: / A:" pairs; headings that read like questions (the heading level
    with the most question-like headings wins); plain paragraphs ending in '?'."""
    bank = Bank(name=name)

    def add(q: str, body: list[str]) -> None:
        q = _clean_question(q)
        if q:
            bank.entries.append(Entry(question=q, response="\n\n".join(b for b in body if b).strip(), bank=name))

    # 1. Q: / A: pairs
    q_pat = re.compile(r"^\s*(?:Q|Question)\s*\d*\s*[:.)\-]\s*", re.I)
    a_pat = re.compile(r"^\s*(?:A|Answer|Response)\s*\d*\s*[:.)\-]\s*", re.I)
    if (sum(1 for lvl, t in blocks if not lvl and q_pat.match(t)) >= 2
            and any(a_pat.match(t) for _, t in blocks)):
        q, body = None, []
        for _, t in blocks:
            if q_pat.match(t):
                if q:
                    add(q, body)
                q, body = q_pat.sub("", t), []
            elif q is not None:
                body.append(a_pat.sub("", t))
        if q:
            add(q, body)
        return bank

    # 2. Headings
    levels = sorted({lvl for lvl, _ in blocks if lvl})
    if levels:
        counts = {lvl: sum(1 for l, t in blocks if l == lvl and _looks_like_question(t)) for lvl in levels}
        best = max(levels, key=lambda l: (counts[l], l))
        if counts[best] == 0:
            best = max(l for l in levels if sum(1 for x, _ in blocks if x == l) >= 1)
        only_questions = counts[best] >= 3  # skip section headings like "Fast facts"
        q, body = None, []
        for lvl, t in blocks:
            if lvl and lvl <= best:
                if q:
                    add(q, body)
                keep = lvl == best and (not only_questions or _looks_like_question(t))
                q, body = (t, []) if keep else (None, [])
            elif q is not None:
                body.append(f"**{t}**" if lvl else t)
        if q:
            add(q, body)
        if len(bank.entries) >= 2:
            return bank
        bank.entries.clear()

    # 3. Paragraphs that end with '?'
    q, body = None, []
    for _, t in blocks:
        if t.strip().endswith("?") and len(t) < 250:
            if q:
                add(q, body)
            q, body = t, []
        elif q is not None:
            body.append(t)
    if q:
        add(q, body)
    return bank


def _split_tables(xml: str) -> list[tuple[str, str]]:
    """Split document XML into top-level ("text", xml) and ("table", xml) parts, respecting nesting."""
    parts, depth, start, last = [], 0, 0, 0
    for m in re.finditer(r"<w:tbl>|</w:tbl>", xml):
        if m.group() == "<w:tbl>":
            if depth == 0:
                parts.append(("text", xml[last:m.start()]))
                start = m.start()
            depth += 1
        else:
            depth -= 1
            if depth == 0:
                parts.append(("table", xml[start:m.end()]))
                last = m.end()
    parts.append(("text", xml[last:]))
    return parts


def _docx_blocks(path: Path) -> tuple[list[tuple[int | None, str]], list[list[list[str]]]]:
    """Read a .docx without extra dependencies: paragraphs (with heading level) and tables.
    Keeps **bold** runs and turns highlighted runs into {{placeholders}}."""
    import html
    import zipfile
    xml = zipfile.ZipFile(path).read("word/document.xml").decode("utf-8")

    def para_text(p: str) -> str:
        out = ""
        for run in re.findall(r"<w:r(?: [^>]*)?>(.*?)</w:r>", p, re.S):
            t = html.unescape("".join(re.findall(r"<w:t(?: [^>]*)?>(.*?)</w:t>", run, re.S)))
            if not t:
                continue
            if "<w:highlight" in run:
                t = "{{" + t + "}}"
            if re.search(r'<w:b/>|<w:b w:val="(?:1|true)"/>', run) and t.strip():
                t = f"**{t}**"
            out += t
        return out.replace("****", "").strip()

    def level(p: str) -> int | None:
        m = re.search(r'<w:pStyle w:val="(?:Heading|heading|Titre|berschrift)(\d)"', p)
        if m:
            return int(m.group(1))
        m = re.search(r'<w:outlineLvl w:val="(\d)"', p)
        return int(m.group(1)) + 1 if m else None

    blocks, tables = [], []
    for kind, part in _split_tables(xml):
        if kind == "table":
            rows = [[" ".join(para_text(p) for p in re.findall(r"<w:p(?: [^>]*)?>.*?</w:p>", c, re.S)).strip()
                     for c in re.findall(r"<w:tc>.*?</w:tc>", r, re.S)]
                    for r in re.findall(r"<w:tr(?: [^>]*)?>.*?</w:tr>", part, re.S)]
            tables.append(rows)
            for row in rows:  # boxed text and small tables still belong to the answer
                cells = [c for c in row if c]
                if cells:
                    blocks.append((None, " · ".join(cells)))
            continue
        for p in re.findall(r"<w:p(?: [^>]*)?>.*?</w:p>", part, re.S):
            if '<w:pStyle w:val="TOC' in p:
                continue
            t = para_text(p)
            if t:
                blocks.append((level(p), t))
    return blocks, tables


def _text_blocks(text: str) -> list[tuple[int | None, str]]:
    blocks = []
    for line in text.splitlines():
        line = line.strip()
        if line:
            m = re.match(r"^(#{1,6}) +(.*)", line)
            blocks.append((len(m.group(1)), m.group(2).strip()) if m else (None, line))
    return blocks


def import_file(path: Path, name: str) -> Bank:
    """Build a bank from a Word doc, spreadsheet, CSV, Markdown or text file.

    Spreadsheets/CSV: a 'question' column and a 'response'/'answer' column (or just the first
    two columns). Documents: questions as headings, "Q:/A:" pairs, or lines ending in '?',
    with the answer in the paragraphs underneath."""
    path = Path(path)
    ext = path.suffix.lower()
    if ext in (".md", ".txt"):
        text = path.read_text(encoding="utf-8-sig")
        if re.search(r"^(also|follows|tags|skeleton):", text, re.M | re.I):
            return parse_markdown(text, name)  # this app's own format, with metadata
        return _blocks_to_bank(_text_blocks(text), name)
    if ext == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as f:
            return _rows_to_bank(list(csv.reader(f)), name)
    if ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        ws = load_workbook(path, read_only=True, data_only=True).active
        return _rows_to_bank([list(r) for r in ws.iter_rows(values_only=True)], name)
    if ext == ".docx":
        blocks, tables = _docx_blocks(path)
        for rows in tables:  # a Q/A table with a recognizable header wins
            if rows and any(c.lower() in COLUMN_NAMES["question"] for c in rows[0]):
                bank = _rows_to_bank(rows, name)
                if bank.entries:
                    return bank
        return _blocks_to_bank(blocks, name)
    raise ValueError(f"Can't import {ext or 'that'} files. Use one of: {', '.join(IMPORT_TYPES)}")


# ---------- config + manager ----------

def load_config() -> dict:
    cfg = dict(DEFAULT_CONFIG)
    if CONFIG_PATH.exists():
        cfg.update(json.loads(CONFIG_PATH.read_text(encoding="utf-8")))
    return cfg


def save_config(cfg: dict) -> None:
    _write_atomic(CONFIG_PATH, json.dumps(cfg, indent=2))


def slug(name: str) -> str:
    s = re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")
    if not s:
        raise ValueError("Bank name needs at least one letter or number.")
    return s


def bank_names() -> list[str]:
    BANKS_DIR.mkdir(exist_ok=True)
    return sorted(p.stem for p in BANKS_DIR.glob("*.md"))


def load_bank(name: str) -> Bank:
    path = BANKS_DIR / f"{name}.md"
    if not path.exists():
        raise KeyError(f"No bank named '{name}'. Existing: {', '.join(bank_names()) or '(none)'}")
    return parse_markdown(path.read_text(encoding="utf-8"), name)


def teach_phrasing(entry: Entry, text: str) -> bool:
    """Save `text` as another way to ask `entry`'s question. Returns False if it was already known."""
    text = " ".join(text.split())
    if not text:
        return False
    bk = load_bank(entry.bank)
    for e in bk.entries:
        if e.question == entry.question:
            known = {p.strip().lower() for p in e.phrasings()}
            if text.lower() in known:
                return False
            e.also.append(text)
            entry.also = list(e.also)   # keep the in-memory entry in step with the file
            bk.save()
            return True
    raise KeyError(f"'{entry.question}' is no longer in bank '{entry.bank}'.")


def active_entries(cfg: dict | None = None) -> list[Entry]:
    cfg = cfg or load_config()
    names = [n for n in cfg["active_banks"] if n in bank_names()]
    return [e for n in names for e in load_bank(n).entries]


DELETED_DIR = BANKS_DIR / "deleted"


def add_bank(name: str, source: Path | None = None, title: str = "", activate: bool = False) -> Bank:
    title = title.strip() or name.strip()
    name = slug(name)
    if name in bank_names():
        raise ValueError(f"A bank named '{name}' already exists. Pick another name.")
    bank = import_file(source, name) if source else Bank(name=name)
    bank.title = title
    bank.save()
    cfg = load_config()
    if activate or not cfg["active_banks"]:
        set_active(name, True)
    return bank


def remove_bank(name: str) -> Path:
    """Move a bank to banks/deleted/ (restorable) and switch it off."""
    bank = load_bank(name)
    DELETED_DIR.mkdir(parents=True, exist_ok=True)
    dest = DELETED_DIR / bank.path.name
    n = 2
    while dest.exists():
        dest = DELETED_DIR / f"{name}-{n}.md"
        n += 1
    bank.path.replace(dest)
    set_active(name, False)
    return dest


def deleted_banks() -> list[Path]:
    return sorted(DELETED_DIR.glob("*.md"), key=lambda p: p.stat().st_mtime, reverse=True) if DELETED_DIR.exists() else []


def restore_bank(path: Path) -> str:
    name = path.stem
    base, n = name, 2
    while name in bank_names():
        name = f"{base}-{n}"
        n += 1
    path.replace(BANKS_DIR / f"{name}.md")
    return name


def rename_bank(name: str, new_title: str) -> str:
    """Change a bank's display title and file name; keeps it active if it was."""
    bank = load_bank(name)
    new = slug(new_title)
    if new != name and new in bank_names():
        raise ValueError(f"A bank named '{new}' already exists.")
    was_active = name in load_config()["active_banks"]
    old_path = bank.path
    bank.name, bank.title = new, new_title.strip()
    bank.save()
    if new != name:
        old_path.unlink()
        set_active(name, False)
        if was_active:
            set_active(new, True)
    return new


def set_active(name: str, on: bool) -> None:
    cfg = load_config()
    active = [n for n in cfg["active_banks"] if n != name]
    cfg["active_banks"] = active + [name] if on else active
    save_config(cfg)


def use_banks(names: list[str]) -> None:
    missing = [n for n in names if n not in bank_names()]
    if missing:
        raise KeyError(f"Unknown bank(s): {', '.join(missing)}")
    cfg = load_config()
    cfg["active_banks"] = names
    save_config(cfg)
